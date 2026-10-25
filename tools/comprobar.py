#!/usr/bin/env python3
"""Pasa documentos por todas las comprobaciones de docs/ingenieria/08_PLAN_DE_PRUEBAS.md § 4.4 y resume.

    python3 tools/comprobar.py fichero.xlsx otro.odt ...
    python3 tools/comprobar.py --ampliado tests/ficheros/calc/valores.ods   # ODF de LibreOffice (con sus extensiones)

Para cada fichero:
  zip         el paquete está bien (no para .fods/.fodt)
  xml         cada parte XML está bien formada
  validador   ODF Validator (estricto; --ampliado: con las extensiones de LibreOffice) o OOXML-Validator
  libreoffice LibreOffice lo abre; su volcado (tools/volcar_lo.py) es la referencia de lo que sigue
  onlyoffice  OnlyOffice (x2t) lo pasa al otro formato y LibreOffice lee eso con los mismos valores y textos
  python      openpyxl o python-docx lo leen con los mismos valores y textos (solo .xlsx y .docx)

Las herramientas, en ~/.cache/lorraine-tools/ (tools/get_herramientas.sh); x2t, el de OnlyOffice del sistema.
Sale con 1 si alguna comprobación falla. -v enseña las primeras diferencias.
"""

import argparse
import json
import math
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.dom import minidom

sys.path.insert(0, str(Path(__file__).resolve().parent))
import volcar_lo  # noqa: E402

TOOLS = Path(os.environ.get("LORRAINE_TOOLS", Path.home() / ".cache" / "lorraine-tools"))
JAR = TOOLS / "odfvalidator-0.13.0-jar-with-dependencies.jar"
OOXML = TOOLS / "ooxml-validator" / "OOXMLValidatorCLI"
VENV_PY = TOOLS / "venv" / "bin" / "python"
X2T = Path("/opt/onlyoffice/desktopeditors/converter/x2t")
OTRO_FORMATO = {".xlsx": ".ods", ".docx": ".odt", ".ods": ".xlsx", ".fods": ".xlsx", ".odt": ".docx", ".fodt": ".docx"}


def comprobar_zip(f):
    if f.suffix in (".fods", ".fodt"):
        return None, "XML plano"
    try:
        with zipfile.ZipFile(f) as z:
            malo = z.testzip()
        return (malo is None), (f"entrada dañada: {malo}" if malo else "bien")
    except zipfile.BadZipFile as e:
        return False, str(e)


def comprobar_xml(f):
    partes = []
    if f.suffix in (".fods", ".fodt"):
        partes.append((f.name, f.read_bytes()))
    else:
        with zipfile.ZipFile(f) as z:
            partes = [(n, z.read(n)) for n in z.namelist() if n.endswith((".xml", ".rels"))]
    for nombre, datos in partes:
        try:
            minidom.parseString(datos)
        except Exception as e:  # noqa: BLE001 (cualquier fallo del analizador es un XML mal formado)
            return False, f"{nombre}: {e}"
    return True, f"{len(partes)} partes"


def comprobar_validador(f, ampliado):
    if f.suffix in (".xlsx", ".docx"):
        r = subprocess.run([str(OOXML), str(f.resolve())], capture_output=True, text=True, timeout=300)
        try:
            errores = json.loads(r.stdout or "[]")
        except json.JSONDecodeError:
            return False, "OOXML-Validator no dio JSON: " + (r.stdout + r.stderr)[:200]
        detalle = "; ".join(sorted({e.get("Description", "")[:90] for e in errores})[:3])
        return not errores, f"{len(errores)} errores" + (f" ({detalle})" if errores else "")
    cmd = ["java", "-jar", str(JAR)] + (["-e"] if ampliado else []) + [str(f)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    errores = [l for l in (r.stdout + r.stderr).splitlines() if "Error:" in l]
    detalle = "; ".join(l.split("Error:", 1)[1].strip()[:90] for l in errores[:3])
    return not errores, f"{len(errores)} errores" + (f" ({detalle})" if errores else "")


def numero(s):
    try:
        return float(s)
    except ValueError:
        return None


def iguales(a, b):
    """Valores del volcado: números con la tolerancia de 08 § 4.3 (un .xlsx guarda 15 cifras), textos tal cual."""
    x, y = numero(a), numero(b)
    if x is not None and y is not None:
        return x == y or math.isclose(x, y, rel_tol=1e-12, abs_tol=1e-300)
    return a == b


def sin_escapes(s):
    """Un campo del volcado tal cual era (volcar_lo.esc al revés)."""
    salida, i = [], 0
    while i < len(s):
        if s[i] == "\\" and i + 1 < len(s):
            salida.append({"t": "\t", "n": "\n", "r": "\r"}.get(s[i + 1], s[i + 1]))
            i += 2
        else:
            salida.append(s[i])
            i += 1
    return "".join(salida)


def contenido(lineas):
    """Lo comparable de un volcado: valores de celdas (sin volátiles) y textos de párrafos y tablas, en orden."""
    salida = []
    for l in lineas:
        c = l.split("\t")
        if c[0] == "celda":
            if any(v in c[5].upper() for v in ("RAND(", "NOW(", "TODAY(")):
                continue
            # un lógico de un .xlsx vuelve como =TRUE(): el tipo y el valor bastan
            salida.append(("celda", c[1], c[2], c[3]))
        elif c[0] in ("parrafo", "fila"):
            salida.append((c[0], c[-1]))
    return salida


def diferencias(ref, otro):
    a, b = contenido(ref), contenido(otro)
    difs = []
    for i in range(max(len(a), len(b))):
        x = a[i] if i < len(a) else None
        y = b[i] if i < len(b) else None
        if x == y:
            continue
        if x and y and x[0] == y[0] == "celda" and x[1:3] == y[1:3] and iguales(x[3], y[3]):
            continue
        # OnlyOffice pasa los errores de celda a texto al convertir: con el mismo texto, vale
        if x and y and x[0] == y[0] == "celda" and x[1] == y[1] and {x[2], y[2]} == {"err", "txt"} and x[3] == y[3]:
            continue
        difs.append(f"{x} | {y}")
    return difs


def comprobar_onlyoffice(f, ref):
    if not X2T.exists():
        return None, "sin OnlyOffice"
    with tempfile.TemporaryDirectory() as d:
        salida = Path(d) / (f.stem + OTRO_FORMATO[f.suffix])
        r = subprocess.run([str(X2T), str(f.resolve()), str(salida)], capture_output=True, text=True,
                           timeout=300, cwd=X2T.parent)
        if r.returncode != 0 or not salida.exists() or salida.stat().st_size == 0:
            return False, f"x2t falló (rc {r.returncode})"
        otro = volcar_lo.volcar([str(salida)])[str(salida)]
    if otro is None:
        return False, "LibreOffice no abre lo que convirtió OnlyOffice"
    difs = diferencias(ref, otro)
    return not difs, (f"{len(difs)} diferencias" if difs else "mismo contenido"), difs


PY_XLSX = r'''
import sys, json, openpyxl
from openpyxl.utils.datetime import to_excel
from openpyxl.utils import get_column_letter
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
out = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                continue
            ref = f"{ws.title}!{get_column_letter(c.column)}{c.row}"
            if isinstance(v, bool):
                out.append([ref, "num", repr(float(v))])
            elif isinstance(v, (int, float)):
                out.append([ref, "num", repr(float(v))])
            elif hasattr(v, "year") or hasattr(v, "hour") or hasattr(v, "days"):
                out.append([ref, "num", repr(float(to_excel(v, wb.epoch)))])
            elif isinstance(v, str) and v.startswith("#") and c.data_type == "e":
                out.append([ref, "err", v])
            else:
                out.append([ref, "txt", str(v)])
print(json.dumps(out))
'''

PY_DOCX = r'''
import sys, json, docx
d = docx.Document(sys.argv[1])
print(json.dumps([p.text for p in d.paragraphs]))
'''


def comprobar_python(f, ref):
    if f.suffix not in (".xlsx", ".docx"):
        return None, "solo .xlsx y .docx"
    if not VENV_PY.exists():
        return None, "sin venv (tools/get_herramientas.sh)"
    codigo = PY_XLSX if f.suffix == ".xlsx" else PY_DOCX
    try:
        r = subprocess.run([str(VENV_PY), "-I", "-c", codigo, str(f.resolve())], capture_output=True, text=True,
                           timeout=300)
    except subprocess.TimeoutExpired:
        return False, "más de 5 minutos (openpyxl recorre todas las filas: limites.xlsx)"
    if r.returncode != 0:
        return False, "no lo lee: " + r.stderr.strip().splitlines()[-1][:120]
    datos = json.loads(r.stdout)
    if f.suffix == ".xlsx":
        celdas = {}
        for l in ref:
            c = l.split("\t")
            if c[0] == "celda" and not any(v in c[5].upper() for v in ("RAND(", "NOW(", "TODAY(")):
                celdas[c[1]] = (c[2], sin_escapes(c[3]))
        leidas = {ref_: (t, v) for ref_, t, v in datos}
        difs = []
        for ref_, (t, v) in celdas.items():
            otro = leidas.get(ref_)
            if otro is None:
                difs.append(f"{ref_}: no está")
            elif otro[0] != t or not iguales(otro[1], v):
                difs.append(f"{ref_}: LibreOffice {t} {v!r}, openpyxl {otro[0]} {otro[1]!r}")
        sobran = [r_ for r_ in leidas if r_ not in celdas]
        difs += [f"{r_}: solo en openpyxl" for r_ in sobran]
    else:
        textos = [sin_escapes(l.split("\t")[-1]) for l in ref if l.startswith("parrafo\t")]
        difs = [f"párrafo {i + 1}: {a!r} | {b!r}" for i, (a, b) in
                enumerate(zip(textos + [None] * len(datos), datos + [None] * len(textos))) if a != b and (a or b)]
    return not difs, (f"{len(difs)} diferencias" if difs else "mismo contenido"), difs


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ficheros", nargs="+")
    ap.add_argument("--ampliado", action="store_true", help="ODF con extensiones de LibreOffice (validador -e)")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()
    if not JAR.exists() or not OOXML.exists():
        sys.exit("faltan las herramientas: tools/get_herramientas.sh")
    rutas = [Path(f) for f in args.ficheros]
    referencias = volcar_lo.volcar([str(f) for f in rutas])
    fallos = 0
    for f in rutas:
        print(f"{f}:")
        ref = referencias.get(str(f))
        pruebas = [("zip", comprobar_zip(f)), ("xml", comprobar_xml(f)),
                   ("validador", comprobar_validador(f, args.ampliado)),
                   ("libreoffice", (ref is not None, f"{len(ref)} líneas de volcado" if ref else "no lo abre"))]
        if ref is not None:
            pruebas += [("onlyoffice", comprobar_onlyoffice(f, ref)), ("python", comprobar_python(f, ref))]
        for nombre, resultado in pruebas:
            ok, texto = resultado[0], resultado[1]
            estado = "  --  " if ok is None else (" bien " if ok else "FALLO ")
            fallos += ok is False
            print(f"  {estado} {nombre:<12} {texto}")
            if args.verbose and ok is False and len(resultado) > 2:
                for d in resultado[2][:8]:
                    print(f"           {d}")
    print(f"\n{fallos} comprobaciones fallidas")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
